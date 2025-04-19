import time
import unittest

import ddddocr
import docx
import openpyxl
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

import utils

MAX_TRIES = 4

# cfg = utils.load_cfg()
wb = openpyxl.Workbook()
# 创建一个sheet并加上名称和所在位置，第一个位置索引号是0
wb.create_sheet("新疆维吾尔自治区农作物病虫疫情监测信息调度管理平台", 0)
sheet = wb["新疆维吾尔自治区农作物病虫疫情监测信息调度管理平台"]
# 写入表头
headers = ["页面", "检测结果"]
sheet.append(headers)

# 创建Word文档对象
doc = docx.Document()


class PPMSXJ():

    def __init__(self):
        # self.cookies = {}
        self.init_chrome()

    def init_chrome(self):
        self.driver = utils.new_chrome()
        self.driver.maximize_window()

    def re_connet(self):
        self.driver.quit()
        self.driver = utils.new_chrome()
        self.driver.maximize_window()
        print("系统重启")

    @utils.retry(MAX_TRIES)
    def test_shouye(self):
        username = input("新疆省级系统生产环境巡检开始\n请输入登录用户名：")
        password = input("请输入登录密码：")
        self.driver.maximize_window()
        self.driver.get("https://xj.pestiot.com:8081")
        self.driver.implicitly_wait(5)

        # 登录
        self.driver.find_element(By.XPATH, '//input[@placeholder="请输入用户名"]').send_keys(username)
        self.driver.find_element(By.XPATH, '//input[@name="password"]').send_keys(password)
        yanzhengma_image = self.driver.find_element(By.XPATH, '//img[contains(@src,"data:image/png;base64")]')
        img_bytes = yanzhengma_image.screenshot_as_png
        yzm = ddddocr.DdddOcr(show_ad=False).classification(img_bytes)
        self.driver.find_element(By.XPATH, '//input[@placeholder="请输入验证码"]').send_keys(yzm)
        time.sleep(3)
        self.driver.find_element(By.XPATH, '//button[@type="button"]').click()

        # 首页元素校验
        try:
            # 等待页面加载完成
            time.sleep(15)
            # 获取页面总高度和宽度
            total_height = self.driver.execute_script("return document.body.parentElement.scrollHeight")
            total_width = self.driver.execute_script("return document.body.parentElement.scrollWidth")
            # 调整窗口尺寸
            self.driver.set_window_size(total_width, total_height)
            element = WebDriverWait(self.driver, 10).until(EC.presence_of_element_located(
                (By.XPATH, '//div[@class="legendBox"]')))
            unittest.TestCase.assertTrue(element is not None, "登录成功，成功打开首页，且指定元素存在")
            utils.g_logger.info(f"登录成功，成功打开首页")
            sheet.append(["首页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/首页.png", doc, "首页")
        except Exception as e:
            utils.g_logger.info("登录成功，但首页显示异常")
            sheet.append(["首页", "异常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/首页.png", doc, "首页")

    @utils.retry(MAX_TRIES)
    def test_shujutianbao(self):
        # 数据填报任务填报页
        try:
            self.driver.find_element(By.XPATH, "//*[text()='数据填报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(EC.visibility_of_element_located(
                (By.XPATH, "//label[text()='报表']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据填报任务填报页，且指定元素存在")
            utils.g_logger.info("数据填报-任务填报页显示正常")
            sheet.append(["数据填报-任务填报", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_任务填报.png", doc, "数据填报_任务填报")
        except Exception as e:
            utils.g_logger.info("数据填报-任务填报页显示异常")
            sheet.append(["数据填报-任务填报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_工作平台.png", doc, "数据填报_任务填报")

        # 数据填报数据查询页
        self.driver.find_element(By.XPATH, "//span[text()='数据查询']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//label[text()='调查时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据填报-数据查询页，且站点列表存在")
            utils.g_logger.info("数据填报-数据查询页显示正常")
            sheet.append(["数据填报-数据查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_数据查询.png", doc, "数据填报_数据查询")
        except Exception as e:
            utils.g_logger.info("数据填报-数据查询页显示异常")
            sheet.append(["数据填报-数据查询", "异常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_数据查询.png", doc, "数据填报_数据查询")

        # 数据填报数据汇总页
        self.driver.find_element(By.XPATH, "//span[text()='数据汇总']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//div[text()=' 显示字段 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据汇总页")
            utils.g_logger.info("数据填报-数据汇总页显示正常")
            sheet.append(["数据填报-数据汇总", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_数据汇总.png", doc, "数据填报_数据汇总")

        except Exception as e:
            utils.g_logger.info("数据填报-数据汇总页显示异常")
            sheet.append(["数据填报-数据汇总", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_数据汇总.png", doc, "数据填报_数据汇总")

        # 数据填报催报查询页
        self.driver.find_element(By.XPATH, "//span[text()='催报信息查询']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//label[text()='站点']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开催报查询页")
            utils.g_logger.info("数据填报-催报信息查询页显示正常")
            sheet.append(["数据填报-催报信息查询", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_催报信息查询.png", doc,
                                  "数据填报_催报信息查询")
        except Exception as e:
            utils.g_logger.info("数据填报-催报信息查询页显示异常")
            sheet.append(["数据填报-催报信息查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_催报信息查询.png", doc,
                                  "数据填报_催报信息查询")

        # 数据填报任务浏览页
        self.driver.find_element(By.XPATH, "//span[text()='任务浏览']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//li[text()='按站点浏览']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开任务浏览页")
            utils.g_logger.info("数据填报-任务浏览页显示正常")
            sheet.append(["数据填报-任务浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_任务浏览.png", doc, "数据填报_任务浏览")
        except Exception as e:
            utils.g_logger.info("数据填报-任务浏览页显示异常")
            sheet.append(["数据填报-任务浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_任务浏览.png", doc, "数据填报_任务浏览")

        # 数据填报报送评价页
        self.driver.find_element(By.XPATH, "//span[text()='报送评价']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//li[text()='按报表统计']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开报送评价页")
            utils.g_logger.info("数据填报-报送评价页显示正常")
            sheet.append(["数据填报-报送评价", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_报送评价.png", doc, "数据填报_报送评价")
        except Exception as e:
            utils.g_logger.info("数据填报-报送评价页显示异常")
            sheet.append(["数据填报-报送评价", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_报送评价.png", doc, "数据填报_报送评价")

        # 数据填报填报任务设置页
        self.driver.find_element(By.XPATH, "//span[text()='任务设置']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='任务起止时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据填报-任务设置页")
            utils.g_logger.info("数据填报-任务设置页显示正常")
            sheet.append(["数据填报-任务设置", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_任务设置.png", doc, "数据填报_任务设置")
        except Exception as e:
            utils.g_logger.info("数据填报-任务设置页异常")
            sheet.append(["数据填报-任务设置", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_任务设置.png", doc,
                                  "数据填报_任务设置")

        # 数据填报农药备品调查页
        self.driver.find_element(By.XPATH, "//span[text()='农药备品调查']").click()
        try:
            time.sleep(10)
            self.driver.find_element(By.XPATH, "//input[@placeholder='请选择']")
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='农药备品数量调查表']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据填报-农药备品调查页")
            utils.g_logger.info("数据填报-农药备品调查页显示正常")
            sheet.append(["数据填报-农药备品调查", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_农药备品调查.png", doc,
                                  "数据填报_农药备品调查")
        except Exception as e:
            utils.g_logger.info("数据填报-农药备品调查页异常")
            sheet.append(["数据填报-农药备品调查", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_农药备品调查.png", doc,
                                  "数据填报_农药备品调查")

        # 数据填报防治组织页
        self.driver.find_element(By.XPATH, "//span[text()='防治组织']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='防治组织名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据填报-防治组织页")
            utils.g_logger.info("数据填报-防治组织页显示正常")
            sheet.append(["数据填报-防治组织", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_防治组织.png", doc, "数据填报_防治组织")
        except Exception as e:
            utils.g_logger.info("数据填报-防治组织页异常")
            sheet.append(["数据填报-防治组织", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据填报_防治组织.png", doc, "数据填报_防治组织")

        # 执行”关闭所有“页面操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(
                EC.element_to_be_clickable(
                    (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']"))
            )
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception as e:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_shujufenxi(self):
        # 数据分析综合看板页
        try:
            self.driver.find_element(By.XPATH, "//*[text()='数据分析 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.visibility_of_any_elements_located((By.XPATH, "//div[@class='el-card__body']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-综合看板页")
            utils.g_logger.info("数据分析-综合看板页显示正常")
            sheet.append(["数据分析-综合看板", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_综合看板.png", doc, "数据分析_综合看板")
        except Exception as e:
            utils.g_logger.info("数据分析-综合看板页显示异常")
            sheet.append(["数据分析-综合看板", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_综合看板.png", doc, "数据分析_综合看板")

        # 数据分析病虫专题页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='病虫专题']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.visibility_of_any_elements_located((By.XPATH, "//div[@class='el-card__body']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-病虫专题页")
            utils.g_logger.info("数据分析-病虫专题页显示正常")
            sheet.append(["数据分析-病虫专题", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫专题.png", doc, "数据分析_病虫专题")
        except Exception as e:
            utils.g_logger.info("数据分析-病虫专题页显示异常")
            sheet.append(["数据分析-病虫专题", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫专题.png", doc, "数据分析_病虫专题")

        # 数据分析GIS在线作图页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='GIS在线作图']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='插值图']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-GIS在线作图页")
            utils.g_logger.info("数据分析-GIS在线作图页显示正常")
            sheet.append(["数据分析-GIS在线作图", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_GIS在线作图.png", doc,
                                  "数据分析_GIS在线作图")
        except Exception as e:
            utils.g_logger.info("数据分析-GIS在线作图页显示异常")
            sheet.append(["数据分析-GIS在线作图", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_GIS在线作图.png", doc,
                                  "数据分析_GIS在线作图")

        # 数据分析病虫趋势页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='病虫趋势']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='是否预测']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-病虫趋势页")
            utils.g_logger.info("数据分析-病虫趋势页显示正常")
            sheet.append(["数据分析-病虫趋势", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫趋势.png", doc, "数据分析_病虫趋势")
        except Exception as e:
            utils.g_logger.info("数据分析-病虫趋势页显示异常")
            sheet.append(["数据分析-病虫趋势", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫趋势.png", doc, "数据分析_病虫趋势")

        # 数据分析自定义分析页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='自定义分析']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='分析类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-自定义分析页")
            utils.g_logger.info("数据分析-自定义分析页显示正常")
            sheet.append(["数据分析-自定义分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_自定义分析.png", doc, "数据分析_自定义分析")
        except Exception as e:
            utils.g_logger.info("数据分析-自定义分析页显示异常")
            sheet.append(["数据分析-自定义分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_自定义分析.png", doc, "数据分析_自定义分析")

        # 数据分析综合看板配置
        try:
            self.driver.find_element(By.XPATH, "//span[text()='综合看板配置']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_all_elements_located((By.XPATH, "//div[@class='el-card__body']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-综合看板配置页")
            utils.g_logger.info("数据分析-综合看板配置页显示正常")
            sheet.append(["数据分析-综合看板配置", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_综合看板配置.png", doc,
                                  "数据分析_综合看板配置")
        except Exception as e:
            utils.g_logger.info("数据分析-综合看板配置页显示异常")
            sheet.append(["数据分析-综合看板配置", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_综合看板配置.png", doc,
                                  "数据分析_综合看板配置")

        # 数据分析病虫专题模型分析
        try:
            self.driver.find_element(By.XPATH, "//span[text()='病虫专题模型分析']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='实时预测']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-病虫专题模型分析页")
            utils.g_logger.info("数据分析-病虫专题模型分析页显示正常")
            sheet.append(["数据分析-病虫专题模型分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫专题模型分析.png", doc,
                                  "数据分析_病虫专题模型分析")
        except Exception as e:
            utils.g_logger.info("数据分析-病虫专题模型分析页显示异常")
            sheet.append(["数据分析-病虫专题模型分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_病虫专题模型分析.png", doc,
                                  "数据分析_病虫专题模型分析")

        # 数据分析数据分析配置
        try:
            self.driver.find_element(By.XPATH, "//span[text()='数据治理']").click()
            self.driver.find_element(By.XPATH, "//li[text()='数据分析配置 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//button/span[text()='新建分析主题']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-数据分析配置页")
            utils.g_logger.info("数据分析-数据分析配置页显示正常")
            sheet.append(["数据分析-数据分析配置", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据分析配置.png", doc,
                                  "数据分析_数据分析配置")
        except Exception as e:
            utils.g_logger.info("数据分析-数据分析配置页显示异常")
            sheet.append(["数据分析-数据分析配置", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据分析配置.png", doc,
                                  "数据分析_数据分析配置")

        # 数据分析数据分析浏览
        try:
            self.driver.find_element(By.XPATH, "//li[text()='数据分析浏览 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//input[@placeholder='按名称搜索']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-数据分析浏览页")
            utils.g_logger.info("数据分析-数据分析浏览页显示正常")
            sheet.append(["数据分析-数据分析浏览", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据分析浏览.png", doc,
                                  "数据分析_数据分析浏览")
        except Exception as e:
            utils.g_logger.info("数据分析-数据分析浏览页显示异常")
            sheet.append(["数据分析-数据分析浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据分析浏览.png", doc,
                                  "数据分析_数据分析浏览")

        # 数据分析数据资源查看
        try:
            self.driver.find_element(By.XPATH, "//li[text()='数据资源查看 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//input[@placeholder='业务范围']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开数据分析-数据资源查看页")
            utils.g_logger.info("数据分析-数据资源查看页显示正常")
            sheet.append(["数据分析-数据资源查看", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据资源查看.png", doc,
                                  "数据分析_数据资源查看")
        except Exception as e:
            utils.g_logger.info("数据分析-数据资源查看页显示异常")
            sheet.append(["数据分析-数据资源查看", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/数据分析_数据资源查看.png", doc,
                                  "数据分析_数据资源查看")
        # 执行”关闭所有“页面操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_wulianwang(self):
        # 物联网数据分析页
        try:
            self.driver.find_element(By.XPATH, "//*[text()='物联网监测 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()=' 物联网统计分析 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-物联网数据分析页")
            utils.g_logger.info("物联网监测-物联网数据分析页显示正常")
            sheet.append(["物联网监测-物联网数据分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网数据分析.png", doc,
                                  "物联网监测-物联网数据分析")
        except Exception as e:
            utils.g_logger.info("物联网监测-物联网数据分析页显示异常")
            sheet.append(["物联网监测-物联网数据分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网数据分析.png", doc,
                                  "物联网监测-物联网数据分析")

        # 物联网监测工作平台页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='工作平台']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='设备分类']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-工作平台页")
            utils.g_logger.info("物联网监测-工作平台页显示正常")
            sheet.append(["物联网监测-工作平台", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-工作平台.png", doc, "物联网监测-工作平台")
        except Exception as e:
            utils.g_logger.info("物联网监测-工作平台页显示异常")
            sheet.append(["物联网监测-工作平台", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-工作平台.png", doc, "物联网监测-工作平台")

        # 物联网监测监测点分布页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='监测点分布']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//button[@title='Zoom in']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-监测点分布页")
            utils.g_logger.info("物联网监测-监测点分布页显示正常")
            sheet.append(["物联网监测-监测点分布", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-监测点分布.png", doc,
                                  "物联网监测-监测点分布")
        except Exception as e:
            utils.g_logger.info("物联网监测-监测点分布页显示异常")
            sheet.append(["物联网监测-监测点分布", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-监测点分布.png", doc,
                                  "物联网监测-监测点分布")

        # 物联网监测设备分布页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='设备分布']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@id='tab-map']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-设备分布页")
            utils.g_logger.info("物联网监测-设备分布页显示正常")
            sheet.append(["物联网监测-设备分布", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-设备分布.png", doc, "物联网监测-设备分布")
        except Exception as e:
            utils.g_logger.info("物联网监测-设备分布页显示异常")
            sheet.append(["物联网监测-设备分布", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-设备分布.png", doc, "物联网监测-设备分布")

        # 环境气象监测-实时数据列表
        try:
            self.driver.find_element(By.XPATH, "//span[text()='环境气象监测']").click()
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='实时数据列表 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='空气温度(°C)']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-环境气象监测-实时数据列表页")
            utils.g_logger.info("物联网监测-环境气象监测-实时数据列表页显示正常")
            sheet.append(["物联网监测-环境气象监测-实时数据列表", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-实时数据列表.png", doc,
                                  "物联网监测-环境气象监测-实时数据列表")
        except Exception as e:
            utils.g_logger.info("物联网监测-环境气象监测-实时数据列表页显示异常")
            sheet.append(["物联网监测-环境气象监测-实时数据列表", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-实时数据列表.png", doc,
                                  "物联网监测-环境气象监测-实时数据列表")

        # 环境气象监测-实时数据推演
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='实时数据推演 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='空气温度']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-环境气象监测-实时数据推演页")
            utils.g_logger.info("物联网监测-环境气象监测-实时数据推演页显示正常")
            sheet.append(["物联网监测-环境气象监测-实时数据推演", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-实时数据推演.png", doc,
                                  "物联网监测-环境气象监测-实时数据推演")
        except Exception as e:
            utils.g_logger.info("物联网监测-环境气象监测-实时数据推演页显示异常")
            sheet.append(["物联网监测-环境气象监测-实时数据推演", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-实时数据推演.png", doc,
                                  "物联网监测-环境气象监测-实时数据推演")

        # 环境气象监测-逐日数据列表
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='逐日数据列表 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='日累计降雨量(mm)']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-环境气象监测-逐日数据列表页")
            utils.g_logger.info("物联网监测-环境气象监测-逐日数据列表页显示正常")
            sheet.append(["物联网监测-环境气象监测-逐日数据列表", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-逐日数据列表.png", doc,
                                  "物联网监测-环境气象监测-逐日数据列表")
        except Exception as e:
            utils.g_logger.info("物联网监测-环境气象监测-逐日数据列表页显示异常")
            sheet.append(["物联网监测-环境气象监测-逐日数据列表", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-逐日数据列表.png", doc,
                                  "物联网监测-环境气象监测-逐日数据列表")

        # 环境气象监测-逐日数据推演
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='逐日数据推演 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='日平均湿度']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-环境气象监测-逐日数据推演页")
            utils.g_logger.info("物联网监测-环境气象监测-逐日数据推演页显示正常")
            sheet.append(["物联网监测-环境气象监测-逐日数据推演", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-逐日数据推演.png", doc,
                                  "物联网监测-环境气象监测-逐日数据推演")
        except Exception as e:
            utils.g_logger.info("物联网监测-环境气象监测-逐日数据推演页显示异常")
            sheet.append(["物联网监测-环境气象监测-逐日数据推演", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-逐日数据推演.png", doc,
                                  "物联网监测-环境气象监测-逐日数据推演")

        # 环境气象监测-趋势分析
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='趋势分析 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='气象指标']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-环境气象监测-趋势分析页")
            utils.g_logger.info("物联网监测-环境气象监测-趋势分析页显示正常")
            sheet.append(["物联网监测-环境气象监测-趋势分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-趋势分析.png", doc,
                                  "物联网监测-环境气象监测-趋势分析")
        except Exception as e:
            utils.g_logger.info("物联网监测-环境气象监测-趋势分析页显示异常")
            sheet.append(["物联网监测-环境气象监测-趋势分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-环境气象监测-趋势分析.png", doc,
                                  "物联网监测-环境气象监测-趋势分析")

        # 性诱监测-逐日数据列表
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='性诱监测']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='逐日数据列表 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='虫量']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-性诱监测-逐日数据列表页")
            utils.g_logger.info("物联网监测-性诱监测-逐日数据列表页显示正常")
            sheet.append(["物联网监测-性诱监测-逐日数据列表", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-逐日数据列表.png", doc,
                                  "物联网监测-性诱监测-逐日数据列表")
        except Exception as e:
            utils.g_logger.info("物联网监测-性诱监测-逐日数据列表页显示异常")
            sheet.append(["物联网监测-性诱监测-逐日数据列表", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-逐日数据列表.png", doc,
                                  "物联网监测-性诱监测-逐日数据列表")

        # 性诱监测-逐日数据推演
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='逐日数据推演 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='可选指标']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-性诱监测-逐日数据推演页")
            utils.g_logger.info("物联网监测-性诱监测-逐日数据推演页显示正常")
            sheet.append(["物联网监测-性诱监测-逐日数据推演", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-逐日数据推演.png", doc,
                                  "物联网监测-性诱监测-逐日数据推演")
        except Exception as e:
            utils.g_logger.info("物联网监测-性诱监测-逐日数据推演页显示异常")
            sheet.append(["物联网监测-性诱监测-逐日数据推演", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-逐日数据推演.png", doc,
                                  "物联网监测-性诱监测-逐日数据推演")

        # 性诱监测-发生趋势分析
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='发生趋势分析 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='统计类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-性诱监测-发生趋势分析页")
            utils.g_logger.info("物联网监测-性诱监测-发生趋势分析页显示正常")
            sheet.append(["物联网监测-性诱监测-发生趋势分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-发生趋势分析.png", doc,
                                  "物联网监测-性诱监测-发生趋势分析")
        except Exception as e:
            utils.g_logger.info("物联网监测-性诱监测-发生趋势分析页显示异常")
            sheet.append(["物联网监测-性诱监测-发生趋势分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-发生趋势分析.png", doc,
                                  "物联网监测-性诱监测-发生趋势分析")

        # 性诱监测-性诱数据统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='性诱数据统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='虫害类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-性诱监测-性诱数据统计页")
            utils.g_logger.info("物联网监测-性诱监测-性诱数据统计页显示正常")
            sheet.append(["物联网监测-性诱监测-性诱数据统计", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-性诱数据统计.png", doc,
                                  "物联网监测-性诱监测-性诱数据统计")
        except Exception as e:
            utils.g_logger.info("物联网监测-性诱监测-性诱数据统计页显示异常")
            sheet.append(["物联网监测-性诱监测-性诱数据统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-性诱监测-性诱数据统计.png", doc,
                                  "物联网监测-性诱监测-性诱数据统计")

        # 灯诱监测-灯诱数据列表
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='灯诱监测']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='灯诱数据列表 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='虫量']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-灯诱监测-灯诱数据列表页")
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据列表页显示正常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据列表", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据列表.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据列表")
        except Exception as e:
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据列表页显示异常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据列表", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据列表.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据列表")

        # 灯诱监测-灯诱数据推演
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='灯诱数据推演 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='可选指标']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-灯诱监测-灯诱数据推演页")
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据推演页显示正常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据推演", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据推演.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据推演")
        except Exception as e:
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据推演页显示异常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据推演", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据推演.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据推演")

        # 灯诱监测-灯诱图片展示
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='灯诱图片展示 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//span[text()='仅看识别后图片']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-灯诱监测-灯诱图片展示页")
            utils.g_logger.info("物联网监测-灯诱监测-灯诱图片展示页显示正常")
            sheet.append(["物联网监测-灯诱监测-灯诱图片展示", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱图片展示.png", doc,
                                  "物联网监测-灯诱监测-灯诱图片展示")
        except Exception as e:
            utils.g_logger.info("物联网监测-灯诱监测-灯诱图片展示页显示异常")
            sheet.append(["物联网监测-灯诱监测-灯诱图片展示", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱图片展示.png", doc,
                                  "物联网监测-灯诱监测-灯诱图片展示")

        # 灯诱监测-发生趋势分析
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='发生趋势分析 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='统计类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-灯诱监测-发生趋势分析页")
            utils.g_logger.info("物联网监测-灯诱监测-发生趋势分析页显示正常")
            sheet.append(["物联网监测-灯诱监测-发生趋势分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-发生趋势分析.png", doc,
                                  "物联网监测-灯诱监测-发生趋势分析")
        except Exception as e:
            utils.g_logger.info("物联网监测-灯诱监测-发生趋势分析页显示异常")
            sheet.append(["物联网监测-灯诱监测-发生趋势分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-发生趋势分析.png", doc,
                                  "物联网监测-灯诱监测-发生趋势分析")

        # 灯诱监测-灯诱数据统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='灯诱数据统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='虫害类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-灯诱监测-灯诱数据统计页")
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据统计页显示正常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据统计", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据统计.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据统计")
        except Exception as e:
            utils.g_logger.info("物联网监测-灯诱监测-灯诱数据统计页显示异常")
            sheet.append(["物联网监测-灯诱监测-灯诱数据统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-灯诱监测-灯诱数据统计.png", doc,
                                  "物联网监测-灯诱监测-灯诱数据统计")

        # 虫量对比分析
        try:
            self.driver.find_element(By.XPATH, "//span[text()='虫量对比分析']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='虫害类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-虫量对比分析页")
            utils.g_logger.info("物联网监测-虫量对比分析页显示正常")
            sheet.append(["物联网监测-虫量对比分析", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-虫量对比分析.png", doc,
                                  "物联网监测-虫量对比分析")
        except Exception as e:
            utils.g_logger.info("物联网监测-虫量对比分析页显示异常")
            sheet.append(["物联网监测-虫量对比分析", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-虫量对比分析.png", doc,
                                  "物联网监测-虫量对比分析")

        # 病害监测-孢子监测
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='病害监测']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='孢子监测 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//*[text()='仅看有孢子']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-病害监测-孢子监测页")
            utils.g_logger.info("物联网监测-病害监测-孢子监测页显示正常")
            sheet.append(["物联网监测-病害监测-孢子监测", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-病害监测-孢子监测.png", doc,
                                  "物联网监测-病害监测-孢子监测")
        except Exception as e:
            utils.g_logger.info("物联网监测-病害监测-孢子监测页显示异常")
            sheet.append(["物联网监测-病害监测-孢子监测", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-病害监测-孢子监测.png", doc,
                                  "物联网监测-病害监测-孢子监测")

        # 视频监控-视频监控分布
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='视频监控']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='视频监控分布 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='站点']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-视频监控-视频监控分布页")
            utils.g_logger.info("物联网监测-视频监控-视频监控分布页显示正常")
            sheet.append(["物联网监测-视频监控-视频监控分布", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-视频监控-视频监控分布.png", doc,
                                  "物联网监测-视频监控-视频监控分布")
        except Exception as e:
            utils.g_logger.info("物联网监测-视频监控-视频监控分布页显示异常")
            sheet.append(["物联网监测-视频监控-视频监控分布", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-视频监控-视频监控分布.png", doc,
                                  "物联网监测-视频监控-视频监控分布")

        # 视频监控-视频图片展示
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='视频图片展示 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_all_elements_located((By.XPATH, "//div[@class='imgCard']/p[text()=' 拍摄 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-视频监控-视频图片展示页")
            utils.g_logger.info("物联网监测-视频监控-视频图片展示页显示正常")
            sheet.append(["物联网监测-视频监控-视频图片展示", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-视频监控-视频图片展示.png", doc,
                                  "物联网监测-视频监控-视频图片展示")
        except Exception as e:
            utils.g_logger.info("物联网监测-视频监控-视频图片展示页显示异常")
            sheet.append(["物联网监测-视频监控-视频图片展示", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-视频监控-视频图片展示.png", doc,
                                  "物联网监测-视频监控-视频图片展示")

        # 物联网管理-设备管理
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='物联网管理']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='设备管理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='设备类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-物联网管理-设备管理页")
            utils.g_logger.info("物联网监测-物联网管理-设备管理页显示正常")
            sheet.append(["物联网监测-物联网管理-设备管理", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网管理-设备管理.png", doc,
                                  "物联网监测-物联网管理-设备管理")
        except Exception as e:
            utils.g_logger.info("物联网监测-物联网管理-设备管理页显示异常")
            sheet.append(["物联网监测-物联网管理-设备管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网管理-设备管理.png", doc,
                                  "物联网监测-物联网管理-设备管理")

        # 物联网管理-监测点管理
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='监测点管理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='监测点名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-物联网管理-监测点管理页")
            utils.g_logger.info("物联网监测-物联网管理-监测点管理页显示正常")
            sheet.append(["物联网监测-物联网管理-监测点管理", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网管理-监测点管理.png", doc,
                                  "物联网监测-物联网管理-监测点管理")
        except Exception as e:
            utils.g_logger.info("物联网监测-物联网管理-监测点管理页显示异常")
            sheet.append(["物联网监测-物联网管理-监测点管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-物联网管理-监测点管理.png", doc,
                                  "物联网监测-物联网管理-监测点管理")

        # 鼠害监测-实时数据列表
        try:
            time.sleep(3)
            self.driver.find_element(By.XPATH, "//span[text()='鼠害监测']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='实时数据列表 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='老鼠种类']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-鼠害监测-实时数据列表")
            utils.g_logger.info("物联网监测-鼠害监测-实时数据列表页显示正常")
            sheet.append(["物联网监测-鼠害监测-实时数据列表", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-鼠害监测-实时数据列表.png", doc,
                                  "物联网监测-鼠害监测-实时数据列表")
        except Exception as e:
            utils.g_logger.info("物联网监测-鼠害监测-实时数据列表页显示异常")
            sheet.append(["物联网监测-鼠害监测-实时数据列表", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-鼠害监测-实时数据列表.png", doc,
                                  "物联网监测-鼠害监测-实时数据列表")

        # 种药处视频
        try:
            self.driver.find_element(By.XPATH, "//span[text()='种药处视频']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//input[@placeholder='请选择']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开物联网监测-种药处视频页")
            utils.g_logger.info("物联网监测-种药处视频页显示正常")
            sheet.append(["物联网监测-种药处视频", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-种药处视频.png", doc,
                                  "物联网监测-种药处视频")
        except Exception as e:
            utils.g_logger.info("物联网监测-种药处视频页显示异常")
            sheet.append(["物联网监测-种药处视频", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/物联网监测-种药处视频.png", doc,
                                  "物联网监测-种药处视频")

        # 执行”关闭所有“页面操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_bingchongyujing(self):
        # 病虫预警页
        try:
            self.driver.find_element(By.XPATH, "//*[text()='病虫预警 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='预警类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开病虫预警页页")
            utils.g_logger.info("病虫预警页显示正常")
            sheet.append(["病虫预警页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫预警页.png", doc,
                                  "病虫预警页")
        except Exception as e:
            utils.g_logger.info("病虫预警页显示异常")
            sheet.append(["病虫预警页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫预警页.png", doc,
                                  "病虫预警页")

        # 病虫情报-情报上传
        try:
            self.driver.find_element(By.XPATH, "//span[text()='病虫情报']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='情报上传 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='发布日期']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开病虫情报-情报上传页")
            utils.g_logger.info("病虫情报-情报上传页显示正常")
            sheet.append(["病虫情报-情报上传", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报上传.png", doc, "病虫情报-情报上传")
        except Exception as e:
            utils.g_logger.info("病虫情报-情报上传页显示异常")
            sheet.append(["病虫情报-情报上传", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报上传.png", doc,
                                  "病虫情报-情报上传")

        # 病虫情报-情报浏览
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='情报浏览 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='期数']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开病虫情报-情报浏览页")
            utils.g_logger.info("病虫情报-情报浏览页显示正常")
            sheet.append(["病虫情报-情报浏览", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报浏览.png", doc, "病虫情报-情报浏览")
        except Exception as e:
            utils.g_logger.info("病虫情报-情报浏览页显示异常")
            sheet.append(["病虫情报-情报浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报浏览.png", doc,
                                  "病虫情报-情报浏览")

        # 病虫情报-情报统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='情报统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@class='cell'][text()='情报类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开病虫情报-情报统计页")
            utils.g_logger.info("病虫情报-情报统计页显示正常")
            sheet.append(["病虫情报-情报统计", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报统计.png", doc, "病虫情报-情报统计")
        except Exception as e:
            utils.g_logger.info("病虫情报-情报统计页显示异常")
            sheet.append(["病虫情报-情报统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/病虫情报-情报统计.png", doc,
                                  "病虫情报-情报统计")

        # 执行“关闭所有页面”操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_zhiwujianyi(self):
        # 植物检疫-产地检疫-产检申请
        try:
            self.driver.find_element(By.XPATH, "//*[text()='植物检疫 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()=' 农业植物产地检疫申请书 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-产检申请页")
            utils.g_logger.info("植物检疫-产地检疫-产检申请页显示正常")
            sheet.append(["植物检疫-产地检疫-产检申请页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检申请.png", doc,
                                  "植物检疫-产地检疫-产检申请")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-产检申请页显示异常")
            sheet.append(["植物检疫-产地检疫-产检申请页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检申请页.png", doc,
                                  "植物检疫-产地检疫-产检申请页")

        # 植物检疫-产地检疫-产检受理
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='产检受理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@id='pane-产检受理']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-产检受理页")
            utils.g_logger.info("植物检疫-产地检疫-产检受理页显示正常")
            sheet.append(["植物检疫-产地检疫-产检受理页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检受理.png", doc,
                                  "植物检疫-产地检疫-产检受理")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-产检受理页显示异常")
            sheet.append(["植物检疫-产地检疫-产检受理页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检受理页.png", doc,
                                  "植物检疫-产地检疫-产检受理页")

        # 植物检疫-产地检疫-田间调查
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='田间调查 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@id='pane-田间调查']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-田间调查页")
            utils.g_logger.info("植物检疫-产地检疫-田间调查页显示正常")
            sheet.append(["植物检疫-产地检疫-田间调查页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-田间调查.png", doc,
                                  "植物检疫-产地检疫-田间调查")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-田间调查页显示异常")
            sheet.append(["植物检疫-产地检疫-田间调查页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-田间调查页.png", doc,
                                  "植物检疫-产地检疫-田间调查页")

        # 植物检疫-产地检疫-实验室检验
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='实验室检验 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='受检单位:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-实验室检验页")
            utils.g_logger.info("植物检疫-产地检疫-实验室检验页显示正常")
            sheet.append(["植物检疫-产地检疫-实验室检验页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-实验室检验.png", doc,
                                  "植物检疫-产地检疫-实验室检验")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-实验室检验页显示异常")
            sheet.append(["植物检疫-产地检疫-实验室检验页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-实验室检验页.png", doc,
                                  "植物检疫-产地检疫-实验室检验页")

        # 植物检疫-产地检疫-签发证书
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='签发证书 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//thead[@class='has-gutter']//div[text()='田调完成时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-签发证书页")
            utils.g_logger.info("植物检疫-产地检疫-签发证书页显示正常")
            sheet.append(["植物检疫-产地检疫-签发证书页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-签发证书.png", doc,
                                  "植物检疫-产地检疫-签发证书")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-签发证书页显示异常")
            sheet.append(["植物检疫-产地检疫-签发证书页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-签发证书页.png", doc,
                                  "植物检疫-产地检疫-签发证书页")

        # 植物检疫-产地检疫-产检查询
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='产检查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//thead[@class='has-gutter']//div[text()='办理状态']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-产检查询页")
            utils.g_logger.info("植物检疫-产地检疫-产检查询页显示正常")
            sheet.append(["植物检疫-产地检疫-产检查询页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检查询.png", doc,
                                  "植物检疫-产地检疫-产检查询")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-产检查询页显示异常")
            sheet.append(["植物检疫-产地检疫-产检查询页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-产检查询页.png", doc,
                                  "植物检疫-产地检疫-产检查询页")

        # 植物检疫-产地检疫-综合查询
        try:
            self.driver.find_element(By.XPATH, "//li[contains(@class,'is-opened')]//li[text()='综合查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='快速查询']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植物检疫-产地检疫-综合查询页")
            utils.g_logger.info("植物检疫-产地检疫-综合查询页显示正常")
            sheet.append(["植物检疫-产地检疫-综合查询页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-综合查询.png", doc,
                                  "植物检疫-产地检疫-综合查询")
        except Exception as e:
            utils.g_logger.info("植物检疫-产地检疫-综合查询页显示异常")
            sheet.append(["植物检疫-产地检疫-综合查询页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植物检疫-产地检疫-综合查询页.png", doc,
                                  "植物检疫-产地检疫-综合查询页")

        # 调运检疫-检疫要求书
        try:
            self.driver.find_element(By.XPATH, "//span[text()='调运检疫']").click()
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='检疫要求书 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//button/span[text()='发检疫要求书']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-检疫要求书页")
            utils.g_logger.info("调运检疫-检疫要求书页显示正常")
            sheet.append(["调运检疫-检疫要求书", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-检疫要求书.png", doc, "调运检疫-检疫要求书")
        except Exception as e:
            utils.g_logger.info("调运检疫-检疫要求书页显示异常")
            sheet.append(["调运检疫-检疫要求书", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-检疫要求书.png", doc,
                                  "调运检疫-检疫要求书")

        # 调检申请-产检换证
        try:
            self.driver.find_element(By.XPATH, "//span[text()='调检申请']").click()
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='产检换证 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@class='el-message-box__header']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-调检申请-产检换证页")
            utils.g_logger.info("调运检疫-调检申请-产检换证页显示正常")
            sheet.append(["调运检疫-调检申请-产检换证", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-产检换证.png", doc,
                                  "调运检疫-调检申请-产检换证")
            self.driver.find_element(By.XPATH, "//i[@class='el-message-box__close el-icon-close']").click()
        except Exception as e:
            utils.g_logger.info("调运检疫-调检申请-产检换证页显示异常")
            sheet.append(["调运检疫-调检申请-产检换证", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-产检换证.png", doc,
                                  "调运检疫-调检申请-产检换证")

        # 调检申请-再次调运
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='再次调运 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//div[@class='el-message-box__header']//span[text()='请输入再调运的植物检疫证书号']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-调检申请-再次调运页")
            utils.g_logger.info("调运检疫-调检申请-再次调运页显示正常")
            sheet.append(["调运检疫-调检申请-再次调运", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-再次调运.png", doc,
                                  "调运检疫-调检申请-再次调运")
            self.driver.find_element(By.XPATH, "//i[@class='el-message-box__close el-icon-close']").click()
        except Exception as e:
            utils.g_logger.info("调运检疫-调检申请-再次调运页显示异常")
            sheet.append(["调运检疫-调检申请-再次调运", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-再次调运.png", doc,
                                  "调运检疫-调检申请-再次调运")

        # 调检申请-直接调运
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='直接调运 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()=' 农业植物调运检疫申请书 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-调检申请-直接调运页")
            utils.g_logger.info("调运检疫-调检申请-直接调运页显示正常")
            sheet.append(["调运检疫-调检申请-直接调运", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-直接调运.png", doc,
                                  "调运检疫-调检申请-直接调运")
        except Exception as e:
            utils.g_logger.info("调运检疫-调检申请-直接调运页显示异常")
            sheet.append(["调运检疫-调检申请-直接调运", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检申请-直接调运.png", doc,
                                  "调运检疫-调检申请-直接调运")

        # 调运检疫-调检受理
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='调检受理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='申请单位或个人']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-调检受理页")
            utils.g_logger.info("调运检疫-调检受理页显示正常")
            sheet.append(["调运检疫-调检受理", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检受理.png", doc, "调运检疫-调检受理")
        except Exception as e:
            utils.g_logger.info("调运检疫-调检受理页显示异常")
            sheet.append(["调运检疫-调检受理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检受理.png", doc,
                                  "调运检疫-调检受理")

        # 调运检疫-检疫检查
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='检疫检查 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='检疫单证号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-检疫检查页")
            utils.g_logger.info("调运检疫-检疫检查页显示正常")
            sheet.append(["调运检疫-检疫检查", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-检疫检查.png", doc, "调运检疫-检疫检查")
        except Exception as e:
            utils.g_logger.info("调运检疫-检疫检查页显示异常")
            sheet.append(["调运检疫-检疫检查", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-检疫检查.png", doc,
                                  "调运检疫-检疫检查")

        # 调运检疫-实验室检验
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='实验室检验 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='检验状态']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-实验室检验页")
            utils.g_logger.info("调运检疫-实验室检验页显示正常")
            sheet.append(["调运检疫-实验室检验", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-实验室检验.png", doc, "调运检疫-实验室检验")
        except Exception as e:
            utils.g_logger.info("调运检疫-实验室检验页显示异常")
            sheet.append(["调运检疫-实验室检验", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-实验室检验.png", doc,
                                  "调运检疫-实验室检验")

        # 调运检疫-签发证书
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='签发证书 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='调运量']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-签发证书")
            utils.g_logger.info("调运检疫-签发证书页显示正常")
            sheet.append(["调运检疫-签发证书", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-签发证书.png", doc, "调运检疫-签发证书")
        except Exception as e:
            utils.g_logger.info("调运检疫-签发证书页显示异常")
            sheet.append(["调运检疫-签发证书", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-签发证书.png", doc,
                                  "调运检疫-签发证书")

        # 调运检疫-调检查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='调检查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='调运类别']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-调检查询")
            utils.g_logger.info("调运检疫-调检查询页显示正常")
            sheet.append(["调运检疫-调检查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检查询.png", doc, "调运检疫-调检查询")
        except Exception as e:
            utils.g_logger.info("调运检疫-调检查询页显示异常")
            sheet.append(["调运检疫-调检查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-调检查询.png", doc,
                                  "调运检疫-调检查询")

        # 调运检疫-综合查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='综合查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='换证查询']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开调运检疫-综合查询")
            utils.g_logger.info("调运检疫-综合查询页显示正常")
            sheet.append(["调运检疫-综合查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-综合查询.png", doc, "调运检疫-综合查询")
        except Exception as e:
            utils.g_logger.info("调运检疫-综合查询页显示异常")
            sheet.append(["调运检疫-综合查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/调运检疫-综合查询.png", doc,
                                  "调运检疫-综合查询")

        # 国外引种-引种申请
        try:
            self.driver.find_element(By.XPATH, "//span[text()='国外引种']").click()
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种申请 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()=' 国（境）外引进农业种苗检疫审批申请书 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种申请页")
            utils.g_logger.info("国外引种-引种申请书页显示正常")
            sheet.append(["国外引种-引种申请", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种申请.png", doc, "国外引种-引种申请")
        except Exception as e:
            utils.g_logger.info("国外引种-引种申请页显示异常")
            sheet.append(["国外引种-引种申请", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种申请.png", doc,
                                  "国外引种-引种申请")

        # 国外引种-引种预审
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种预审 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='引种单位']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种预审页")
            utils.g_logger.info("国外引种-引种预审书页显示正常")
            sheet.append(["国外引种-引种预审", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种预审.png", doc, "国外引种-引种预审")
        except Exception as e:
            utils.g_logger.info("国外引种-引种预审页显示异常")
            sheet.append(["国外引种-引种预审", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种预审.png", doc,
                                  "国外引种-引种预审")

        # 国外引种-引种受理
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种受理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='引种类型：']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种受理页")
            utils.g_logger.info("国外引种-引种受理书页显示正常")
            sheet.append(["国外引种-引种受理", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种受理.png", doc, "国外引种-引种受理")
        except Exception as e:
            utils.g_logger.info("国外引种-引种受理页显示异常")
            sheet.append(["国外引种-引种受理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种受理.png", doc,
                                  "国外引种-引种受理")

        # 国外引种-审批签发
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='审批签发 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='承诺截止时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-审批签发页")
            utils.g_logger.info("国外引种-审批签发书页显示正常")
            sheet.append(["国外引种-审批签发", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-审批签发.png", doc, "国外引种-审批签发")
        except Exception as e:
            utils.g_logger.info("国外引种-审批签发页显示异常")
            sheet.append(["国外引种-审批签发", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-审批签发.png", doc,
                                  "国外引种-审批签发")

        # 国外引种-引种查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='拟种植省份']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种查询页")
            utils.g_logger.info("国外引种-引种查询页显示正常")
            sheet.append(["国外引种-引种查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种查询.png", doc, "国外引种-引种查询")
        except Exception as e:
            utils.g_logger.info("国外引种-引种查询页显示异常")
            sheet.append(["国外引种-引种查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种查询.png", doc,
                                  "国外引种-引种查询")

        # 国外引种-引种跟踪
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种跟踪 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='是否入境登记：']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种跟踪页")
            utils.g_logger.info("国外引种-引种跟踪页显示正常")
            sheet.append(["国外引种-引种跟踪", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种跟踪.png", doc, "国外引种-引种跟踪")
        except Exception as e:
            utils.g_logger.info("国外引种-引种跟踪页显示异常")
            sheet.append(["国外引种-引种跟踪", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种跟踪.png", doc,
                                  "国外引种-引种跟踪")

        # 国外引种-引种田调
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种田调 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='检疫单证号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种田调页")
            utils.g_logger.info("国外引种-引种田调页显示正常")
            sheet.append(["国外引种-引种田调", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种田调.png", doc, "国外引种-引种田调")
        except Exception as e:
            utils.g_logger.info("国外引种-引种田调页显示异常")
            sheet.append(["国外引种-引种田调", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种田调.png", doc,
                                  "国外引种-引种田调")

        # 国外引种-引种检验
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='引种检验 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='送样单位:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-引种检验页")
            utils.g_logger.info("国外引种-引种检验页显示正常")
            sheet.append(["国外引种-引种检验", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种检验.png", doc, "国外引种-引种检验")
        except Exception as e:
            utils.g_logger.info("国外引种-引种检验页显示异常")
            sheet.append(["国外引种-引种检验", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种检验.png", doc,
                                  "国外引种-引种检验")

        # 国外引种-综合查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='综合查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[text()='查询引种许可信息']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开国外引种-综合查询页")
            utils.g_logger.info("国外引种-综合查询页显示正常")
            sheet.append(["国外引种-综合查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-引种检验.png", doc, "国外引种-综合查询")
        except Exception as e:
            utils.g_logger.info("国外引种-综合查询页显示异常")
            sheet.append(["国外引种-综合查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/国外引种-综合查询.png", doc,
                                  "国外引种-综合查询")

        # 疫情报告-快报
        try:
            time.sleep(5)
            self.driver.find_element(By.XPATH, "//span[text()='疫情报告']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='快报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='快报编号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开疫情报告-快报页")
            utils.g_logger.info("疫情报告-快报页显示正常")
            sheet.append(["疫情报告-快报", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-快报.png", doc, "疫情报告-快报")
        except Exception as e:
            utils.g_logger.info("疫情报告-快报页显示异常")
            sheet.append(["疫情报告-快报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-快报.png", doc,
                                  "疫情报告-快报")

        # 疫情报告-月报
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='月报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='月报编号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开疫情报告-月报页")
            utils.g_logger.info("疫情报告-月报页显示正常")
            sheet.append(["疫情报告-月报", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-月报.png", doc, "疫情报告-月报")
        except Exception as e:
            utils.g_logger.info("疫情报告-月报页显示异常")
            sheet.append(["疫情报告-月报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-月报.png", doc,
                                  "疫情报告-月报")

        # 疫情报告-年报
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='年报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='年报编号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开疫情报告-年报页")
            utils.g_logger.info("疫情报告-年报页显示正常")
            sheet.append(["疫情报告-年报", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-年报.png", doc, "疫情报告-年报")
        except Exception as e:
            utils.g_logger.info("疫情报告-年报页显示异常")
            sheet.append(["疫情报告-年报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/疫情报告-年报.png", doc,
                                  "疫情报告-年报")

        # 检疫员-检疫员查询
        try:
            time.sleep(5)
            self.driver.find_element(By.XPATH, "//span[text()='检疫员']").click()
            time.sleep(5)
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='检疫员查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//div[@class='search-container']//label[text()='检疫员编号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开检疫员-检疫员查询页")
            utils.g_logger.info("检疫员-检疫员查询页显示正常")
            sheet.append(["检疫员-检疫员查询", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-检疫员查询.png", doc, "检疫员-检疫员查询")
        except Exception as e:
            utils.g_logger.info("检疫员-检疫员查询页显示异常")
            sheet.append(["检疫员-检疫员查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-检疫员查询.png", doc,
                                  "检疫员-检疫员查询")

        # 检疫员-新增检疫员
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='新增检疫员 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//button/span[text()='新增检疫员']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开检疫员-新增检疫员页")
            utils.g_logger.info("检疫员-新增检疫员页显示正常")
            sheet.append(["检疫员-新增检疫员", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-新增检疫员.png", doc, "检疫员-新增检疫员")
        except Exception as e:
            utils.g_logger.info("检疫员-新增检疫员页显示异常")
            sheet.append(["检疫员-新增检疫员", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-新增检疫员.png", doc,
                                  "检疫员-新增检疫员")

        # 检疫员-检疫员管理
        try:
            self.driver.find_element(By.XPATH,
                                     "//li[contains(@class,'is-opened')]//li[text()='检疫员管理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//div[@class='search-container']//label[text()='检疫员编号:']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开检疫员-检疫员管理页")
            utils.g_logger.info("检疫员-检疫员管理页显示正常")
            sheet.append(["检疫员-检疫员管理", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-检疫员管理.png", doc, "检疫员-检疫员管理")
        except Exception as e:
            utils.g_logger.info("检疫员-检疫员管理页显示异常")
            sheet.append(["检疫员-检疫员管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/检疫员-检疫员管理.png", doc,
                                  "检疫员-检疫员管理")

        # 执行“关闭所有页面”操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_zhihuidiaodu(self):
        # 指挥调度页
        try:
            self.driver.find_element(By.XPATH, "//*[text()='指挥调度 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[contains(@id,'legendBox')]")))
            unittest.TestCase.assertTrue(element is not None, "成功打开指挥调度页")
            utils.g_logger.info("指挥调度页显示正常")
            sheet.append(["指挥调度页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/指挥调度页.png", doc,
                                  "指挥调度页")
        except Exception as e:
            utils.g_logger.info("指挥调度页显示异常")
            sheet.append(["指挥调度页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/指挥调度页.png", doc,
                                  "指挥调度页")

        # 蝗虫指挥调度页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='蝗虫指挥调度']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//div[@class='title'][text()='新疆农区蝗虫发生防控调度']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开蝗虫指挥调度页")
            utils.g_logger.info("蝗虫指挥调度页显示正常")
            sheet.append(["蝗虫指挥调度页", "正常"])
            # 获取网页截图并保存至word文档
            utils.page_screenshot(self.driver, "outputs/imagefiles/蝗虫指挥调度页.png", doc,
                                  "蝗虫指挥调度页")
        except Exception as e:
            utils.g_logger.info("蝗虫指挥调度页显示异常")
            sheet.append(["蝗虫指挥调度页", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/蝗虫指挥调度页.png", doc,
                                  "蝗虫指挥调度页")

        # 执行“关闭所有页面”操作
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_xitongguanli(self):
        # 系统管理机构管理页
        time.sleep(5)
        self.driver.find_element(By.XPATH, "//*[text()='系统管理 ']").click()
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='机构管理']").click()
        try:
            time.sleep(10)
            orgnization_tree_element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//label[text()='机构名称']")))
            unittest.TestCase.assertTrue(orgnization_tree_element is not None, "成功打开机构管理页")
            utils.g_logger.info("系统管理-机构管理页显示正常")
            sheet.append(["系统管理-机构管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/机构管理.png", doc, "机构管理")
        except Exception as e:
            utils.g_logger.info("系统管理-机构管理页异常")
            sheet.append(["系统管理-机构管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/机构管理.png", doc, "机构管理")

        # 系统管理用户管理页
        self.driver.find_element(By.XPATH, "//span[text()='用户管理']").click()
        try:
            time.sleep(10)
            orgnization_tree_element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//label[text()='用户名称']")))
            unittest.TestCase.assertTrue(orgnization_tree_element is not None, "成功打开用户管理页")
            utils.g_logger.info("系统管理-用户管理页显示正常")
            sheet.append(["系统管理-用户管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/用户管理.png", doc, "用户管理")
        except Exception as e:
            utils.g_logger.info("系统管理-用户管理页异常")
            sheet.append(["系统管理-用户管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/用户管理.png", doc, "用户管理")

        # 系统管理角色管理页
        self.driver.find_element(By.XPATH, "//span[text()='角色管理']").click()
        try:
            time.sleep(13)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='角色名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-角色管理页")
            utils.g_logger.info("系统管理-角色管理页显示正常")
            sheet.append(["系统管理-角色管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_角色管理.png", doc, "系统管理_角色管理")
        except Exception as e:
            utils.g_logger.info("系统管理-角色管理页异常")
            sheet.append(["系统管理-角色管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_角色管理.png", doc, "系统管理_角色管理")

        """"# 系统管理菜单管理页
        self.driver.find_element(By.XPATH, "//span[text()='菜单管理']").click()
        try:
            time.sleep(12)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//label[text()='菜单']")))
            unittest.TestCase.assertTrue(element is not None,
                                         "成功打开系统管理-菜单管理页")
            utils.g_logger.info("系统管理-菜单管理页显示正常")
            sheet.append(["系统管理-菜单管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_菜单管理.png", doc, "系统管理_菜单管理")
        except Exception as e:
            utils.g_logger.info("系统管理-菜单管理页异常")
            sheet.append(["系统管理-菜单管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_菜单管理.png", doc, "系统管理_菜单管理")"""

        # 系统管理日志管理登录日志页
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='日志管理']").click()
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='登录日志 ']").click()
        try:
            time.sleep(11)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//label[@class='el-form-item__label'][text()='登录时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-登录日志页")
            utils.g_logger.info("系统管理-登录日志页显示正常")
            sheet.append(["系统管理-登录日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_登录日志.png", doc,
                                  "系统管理_日志管理_登录日志")
        except Exception as e:
            utils.g_logger.info("系统管理-登录日志页异常")
            sheet.append(["系统管理-登录日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_登录日志.png", doc,
                                  "系统管理_日志管理_登录日志")

        # 系统管理日志管理操作日志页
        try:
            self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='操作日志 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//label[@class='el-form-item__label'][text()='操作类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-操作日志页，且指定元素存在")
            utils.g_logger.info("系统管理-操作日志页显示正常")
            sheet.append(["系统管理-操作日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_操作日志.png", doc,
                                  "系统管理_日志管理_操作日志")
        except Exception as e:
            utils.g_logger.info("系统管理-操作日志页异常")
            sheet.append(["系统管理-操作日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_操作日志.png", doc,
                                  "系统管理_日志管理_操作日志")

        """# 系统管理日志管理上报日志页
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='上报日志 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located((By.XPATH, "//label[@class='el-form-item__label'][text()='上报国家状态']")))
            unittest.TestCase.assertTrue(element is not None,"成功打开系统管理-上报日志页")
            utils.g_logger.info("系统管理-上报日志页显示正常")
            sheet.append(["系统管理-上报日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_上报日志.png", doc,"系统管理_日志管理_上报日志")
        except Exception as e:
            utils.g_logger.info("系统管理-上报日志页异常")
            sheet.append(["系统管理-上报日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_上报日志.png", doc,"系统管理_日志管理_上报日志")

        # 系统管理日志管理同步日志页
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='同步日志 ']").click()
        try:
            time.sleep(12)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located((By.XPATH, "//label[@class='el-form-item__label'][text()='同步时间']")))
            unittest.TestCase.assertTrue(element is not None,"成功打开系统管理-同步日志页")
            utils.g_logger.info("系统管理-同步日志页显示正常")
            sheet.append(["系统管理-同步日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_同步日志.png", doc,"系统管理_日志管理_同步日志")
        except Exception as e:
            utils.g_logger.info("系统管理-同步日志页异常")
            sheet.append(["系统管理-同步日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_同步日志.png", doc,"系统管理_日志管理_同步日志")

        # 系统管理日志管理物联网维护日志页
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='物联网维护日志 ']").click()
        try:
            time.sleep(12)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located((By.XPATH, "//label[text()='站点']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-物联网维护日志页")
            utils.g_logger.info("系统管理-物联网维护日志页显示正常")
            sheet.append(["系统管理-物联网维护日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_物联网维护日志.png", doc,
                                  "系统管理_日志管理_物联网维护日志")
        except Exception as e:
            utils.g_logger.info("系统管理-物联网维护日志页异常")
            sheet.append(["系统管理-物联网维护日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_物联网维护日志.png", doc,
                                  "系统管理_日志管理_物联网维护日志")

        # 系统管理日志管理错误日志页
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='错误日志 ']").click()
        try:
            time.sleep(12)
            element = WebDriverWait(self.driver, 5).until(
                EC.presence_of_element_located((By.XPATH, "//div[@class='cell'][text()='错误信息']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-错误日志页")
            utils.g_logger.info("系统管理-错误日志页显示正常")
            sheet.append(["系统管理-错误日志", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_错误日志.png", doc,
                                  "系统管理_日志管理_错误日志")
        except Exception as e:
            utils.g_logger.info("系统管理-错误日志页异常")
            sheet.append(["系统管理-错误日志", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_日志管理_错误日志.png", doc,
                                  "系统管理_日志管理_错误日志")

        # 系统管理字典表管理页
        self.driver.find_element(By.XPATH, "//span[text()='字典表管理']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//h3[text()='字典表列表']")))
            unittest.TestCase.assertTrue(element is not None,"成功打开系统管理-字典表管理页，且指定元素存在")
            utils.g_logger.info("系统管理-字典表管理页显示正常")
            sheet.append(["系统管理-字典表管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_字典表管理.png", doc, "系统管理_字典表管理")
        except Exception as e:
            utils.g_logger.info("系统管理-字典表管理页异常")
            sheet.append(["系统管理-字典表管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_字典表管理.png", doc, "系统管理_字典表管理")"""

        # 系统管理报表权限管理页
        try:
            self.driver.find_element(By.XPATH, "//span[text()='报表权限管理']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, "//thead[@class='has-gutter']//div[text()='可操作表']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开系统管理-报表权限管理页")
            utils.g_logger.info("系统管理-报表权限管理页显示正常")
            sheet.append(["系统管理-报表权限管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_报表权限管理.png", doc,
                                  "系统管理_报表权限管理")
        except Exception as e:
            utils.g_logger.info("系统管理-报表权限管理页异常")
            sheet.append(["系统管理-报表权限管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/系统管理_字典表管理.png", doc,
                                  "系统管理_报表权限管理")

        # 关闭所有页面
        try:
            self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
            time.sleep(3)
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_zhishiku(self):
        # 作物知识库-知识浏览
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//*[text()='知识库 ']").click()
        try:
            time.sleep(6)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@role='radiogroup']//span[text()='粮食作物']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-作物知识库-知识浏览页，且指定元素存在")
            sheet.append(["知识库-作物知识库-知识浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/作物知识库_知识浏览.png", doc, "作物知识库_知识浏览")
        except Exception as e:
            utils.g_logger.info("知识库-作物知识库-知识浏览页异常")
            sheet.append(["知识库-作物知识库-知识浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/作物知识库_知识浏览.png", doc, "作物知识库_知识浏览")

        # 作物知识库-知识维护
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='作物知识库']/parent::div/following-sibling::ul//li[text()='知识维护 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='中文名']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-作物知识库-知识维护页，且指定元素存在")
            utils.g_logger.info("知识库-作物知识库-知识维护页显示正常")
            sheet.append(["知识库-作物知识库-知识维护", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/作物知识库_知识维护.png", doc, "作物知识库_知识维护")
        except Exception as e:
            utils.g_logger.info("知识库-作物知识库-知识维护页异常")
            sheet.append(["知识库-作物知识库-知识维护", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/作物知识库_知识维护.png", doc, "作物知识库_知识维护")

        # 知识库-病虫害知识库-知识浏览页
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='病虫害知识库']").click()
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='病虫害知识库']/parent::div/following-sibling::ul//li[text()='知识浏览 ']").click()
        try:
            time.sleep(10)
            elements = WebDriverWait(self.driver, 10).until(
                EC.presence_of_all_elements_located((By.XPATH, "//div[@class='el-image img']/img")))
            unittest.TestCase.assertTrue(elements is not None, "成功打开知识库-病虫害知识库-知识浏览页")
            utils.g_logger.info("知识库-病虫害知识库-知识浏览页显示正常")
            sheet.append(["知识库-病虫害知识库-知识浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_病虫害知识库_知识浏览.png", doc,
                                  "知识库_病虫害知识库_知识浏览")
        except Exception as e:
            utils.g_logger.info("知识库-病虫害知识库-知识浏览页异常")
            sheet.append(["知识库-病虫害知识库-知识浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_病虫害知识库_知识浏览.png", doc,
                                  "知识库_病虫害知识库_知识浏览")

        # 知识库-病虫害知识库-知识维护页
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='病虫害知识库']/parent::div/following-sibling::ul//li[text()='知识维护 ']").click()
        try:
            time.sleep(9)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='修改时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-病虫害知识库-知识维护页")
            utils.g_logger.info("知识库-病虫害知识库-知识维护页显示正常")
            sheet.append(["知识库-病虫害知识库-知识维护", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_病虫害知识库_知识维护.png", doc,
                                  "知识库_病虫害知识库_知识维护")
        except Exception as e:
            utils.g_logger.info("知识库-病虫害知识库-知识维护页异常")
            sheet.append(["知识库-病虫害知识库-知识维护", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_病虫害知识库_知识维护.png", doc,
                                  "知识库_病虫害知识库_知识维护")

        # 植保知识库-知识浏览页
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='植保知识库']").click()
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='植保知识库']/parent::div/following-sibling::ul//li[text()='知识浏览 ']").click()
        try:
            time.sleep(8)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@role='radiogroup']//span[text()='植物检疫']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-植保知识库-知识浏览页")
            utils.g_logger.info("知识库-植保知识库-知识浏览页显示正常")
            sheet.append(["知识库-植保知识库-知识浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识浏览.png", doc,
                                  "知识库_植保知识库_知识浏览")
        except Exception as e:
            utils.g_logger.info("知识库-植保知识库-知识浏览页异常")
            sheet.append(["知识库-植保知识库-知识浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识浏览.png", doc,
                                  "知识库_植保知识库_知识浏览")

        # 植保知识库-知识审核页
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='植保知识库']/parent::div/following-sibling::ul//li[text()='知识审核 ']").click()
        try:
            time.sleep(8)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='发布时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-植保知识库-知识审核页，且指定元素存在")
            utils.g_logger.info("知识库-植保知识库-知识审核页显示正常")
            sheet.append(["知识库-植保知识库-知识审核", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识审核.png", doc,
                                  "知识库_植保知识库_知识审核")
        except Exception as e:
            utils.g_logger.info("知识库-植保知识库-知识审核页异常")
            sheet.append(["知识库-植保知识库-知识审核", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识审核.png", doc,
                                  "知识库_植保知识库_知识审核")

        # 植知识库-知识上传页
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='植保知识库']/parent::div/following-sibling::ul//li[text()='知识上传 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='发布时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-植保知识库-知识上传页")
            utils.g_logger.info("知识库-植保知识库-知识上传页显示正常")
            sheet.append(["知识库-植保知识库-知识上传", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识上传.png", doc,
                                  "知识库_植保知识库_知识上传")
        except Exception as e:
            utils.g_logger.info("知识库-植保知识库-知识上传页异常")
            sheet.append(["知识库-植保知识库-知识上传", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_植保知识库_知识上传.png", doc,
                                  "知识库_植保知识库_知识上传")

        # 资料库-资料库浏览页
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='资料库']").click()
        time.sleep(3)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='资料库']/parent::div/following-sibling::ul//li[text()='资料库浏览 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//span[@class='el-radio-button__inner'][text()='资料库']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-资料库-资料库浏览页")
            utils.g_logger.info("知识库-资料库-资料库浏览页显示正常")
            sheet.append(["知识库-资料库-资料库浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_资料库_资料库浏览.png", doc,
                                  "知识库_资料库_资料库浏览")
        except Exception as e:
            utils.g_logger.info("知识库-资料库-资料库浏览页异常")
            sheet.append(["知识库-资料库-资料库浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_资料库_资料库浏览.png", doc,
                                  "知识库_资料库_资料库浏览")

        # 资料库-资料库维护
        self.driver.find_element(By.XPATH,
                                 "//span[text()='资料库']/parent::div/following-sibling::ul//li[text()='资料库维护 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//div[@class='search-container']//label[text()='目录']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开知识库-资料库-资料库维护页")
            utils.g_logger.info("知识库-资料库-资料库维护页显示正常")
            sheet.append(["知识库-资料库-资料库维护", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_资料库_资料库维护.png", doc,
                                  "知识库_资料库_资料库维护")
        except Exception as e:
            utils.g_logger.info("知识库-资料库-资料库维护页异常")
            sheet.append(["知识库-资料库-资料库维护", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/知识库_资料库_资料库维护.png", doc,
                                  "知识库_资料库_资料库维护")

        # 关闭所有页面
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_bangongyingyong(self):
        # 办公应用-视频会议-会议管理
        time.sleep(2)
        self.driver.find_element(By.XPATH, "//*[text()='办公应用 ']").click()
        try:
            time.sleep(8)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='参会人员']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-视频会议-会议管理页")
            utils.g_logger.info("办公应用-视频会议-会议管理页显示正常")
            sheet.append(["办公应用-视频会议-会议管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_视频会议_会议管理.png", doc,
                                  "办公应用_视频会议_会议管理")
        except Exception as e:
            utils.g_logger.info("办公应用-视频会议-会议管理页异常")
            sheet.append(["办公应用-视频会议-会议管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_视频会议_会议管理.png", doc,
                                  "办公应用_视频会议_会议管理")

        # 办公应用-视频会议-我的会议
        try:
            self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='我的会议 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//label[@class='el-form-item__label'][text()='会议状态']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-视频会议-我的会议页")
            utils.g_logger.info("办公应用-视频会议-我的会议页显示正常")
            sheet.append(["办公应用-视频会议-我的会议", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_视频会议_我的会议.png", doc,
                                  "办公应用_视频会议_我的会议")
        except Exception as e:
            utils.g_logger.info("办公应用-视频会议-我的会议页异常")
            sheet.append(["办公应用-视频会议-我的会议", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_视频会议_我的会议.png", doc,
                                  "办公应用_视频会议_我的会议")

        # 办公应用-文件收发管理-收件箱页
        time.sleep(3)
        try:
            self.driver.find_element(By.XPATH, "//span[text()='文件收发']").click()
            time.sleep(3)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='文件收发']/parent::div/following-sibling::ul//li[text()='收件箱 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='收件时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-文件收发管理-收件箱页")
            utils.g_logger.info("办公应用-文件收发管理-收件箱页显示正常")
            sheet.append(["办公应用-文件收发管理-收件箱", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_收件箱.png", doc,
                                  "办公应用_文件收发管理_收件箱")
        except Exception as e:
            utils.g_logger.info(e)
            sheet.append(["办公应用-文件收发管理-收件箱", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_收件箱.png", doc,
                                  "办公应用_文件收发管理_收件箱")

        # 办公应用-文件收发管理-草稿箱页
        time.sleep(3)
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='文件收发']/parent::div/following-sibling::ul//li[text()='草稿箱 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='收件单位']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-文件收发管理-草稿箱页")
            utils.g_logger.info("办公应用-文件收发管理-草稿箱页显示正常")
            sheet.append(["办公应用-文件收发管理-草稿箱", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_草稿箱.png", doc,
                                  "办公应用_文件收发管理_草稿箱")
        except Exception as e:
            utils.g_logger.info("办公应用-文件收发管理-草稿箱页异常")
            sheet.append(["办公应用-文件收发管理-草稿箱", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_草稿箱.png", doc,
                                  "办公应用_文件收发管理_草稿箱")

        # 办公应用-文件收发管理-发件箱页
        time.sleep(3)
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='文件收发']/parent::div/following-sibling::ul//li[text()='发件箱 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='发送时间']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-文件收发管理-发件箱页")
            utils.g_logger.info("办公应用-文件收发管理-发件箱页显示正常")
            sheet.append(["办公应用-文件收发管理-发件箱", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_发件箱.png", doc,
                                  "办公应用_文件收发管理_发件箱")
        except Exception as e:
            utils.g_logger.info("办公应用-文件收发管理-发件箱页异常")
            sheet.append(["办公应用-文件收发管理-发件箱", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_文件收发管理_发件箱.png", doc,
                                  "办公应用_文件收发管理_发件箱")

        """# 病虫害情报-情报管理页
        time.sleep(3)
        self.driver.find_element(By.XPATH, "//span[text()='病虫害情报']").click()
        time.sleep(2)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='病虫害情报']/parent::div/following-sibling::ul//li[text()='情报管理 ']").click()
        try:
            time.sleep(15)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//label[text()='情报类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-病虫害情报-情报管理页，且指定元素存在")
            utils.g_logger.info("办公应用-病虫害情报-情报管理页显示正常")
            sheet.append(["办公应用-病虫害情报-情报管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报管理.png", doc,
                                  "办公应用_病虫害情报_情报管理")
        except Exception as e:
            utils.g_logger.info("办公应用-病虫害情报-情报管理页异常")
            sheet.append(["办公应用-病虫害情报-情报管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报管理.png", doc,
                                  "办公应用_病虫害情报_情报管理")

        # 病虫害情报-情报库检索页
        time.sleep(5)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='病虫害情报']/parent::div/following-sibling::ul//li[text()='情报库检索 ']").click()
        try:
            time.sleep(15)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//label[text()='关键词']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-病虫害情报-情报库检索页，且指定元素存在")
            utils.g_logger.info("办公应用-病虫害情报-情报库检索页显示正常")
            sheet.append(["办公应用-病虫害情报-情报库检索", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报库检索.png", doc,
                                  "办公应用_病虫害情报_情报库检索")
        except Exception as e:
            utils.g_logger.info("办公应用-病虫害情报-情报库检索页异常")
            sheet.append(["办公应用-病虫害情报-情报库检索", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报库检索.png", doc,
                                  "办公应用_病虫害情报_情报库检索")

        # 病虫害情报-情报统计页
        time.sleep(5)
        self.driver.find_element(By.XPATH,
                                 "//span[text()='病虫害情报']/parent::div/following-sibling::ul//li[text()='情报统计 ']").click()
        try:
            time.sleep(10)
            elements = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_any_elements_located(
                    (By.XPATH, "//span[@class='el-radio-button__inner'][contains(text(),'统计')]")))
            unittest.TestCase.assertTrue(elements is not None, "成功打开办公应用-病虫害情报-情报统计页，且指定元素存在")
            utils.g_logger.info("办公应用-病虫害情报-情报统计页显示正常")
            sheet.append(["办公应用-病虫害情报-情报统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报统计.png", doc,
                                  "办公应用_病虫害情报_情报统计")
        except Exception as e:
            utils.g_logger.info("办公应用-病虫害情报-情报统计页异常")
            sheet.append(["办公应用-病虫害情报-情报统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_病虫害情报_情报统计.png", doc,
                                  "办公应用_病虫害情报_情报统计")

        # 新闻管理-新闻浏览页
        time.sleep(5)
        self.driver.find_element(By.XPATH, "//span[text()='新闻管理']/following-sibling::i").click()
        time.sleep(2)
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='新闻浏览 ']").click()
        try:
            time.sleep(10)
            elements = WebDriverWait(self.driver, 10).until(
                EC.presence_of_all_elements_located((By.XPATH, "//h2[@class='zbtitle']")))
            unittest.TestCase.assertTrue(elements is not None, "成功打开办公应用-新闻管理-新闻浏览页，且指定元素存在")
            utils.g_logger.info("办公应用-新闻管理-新闻浏览页显示正常")
            sheet.append(["办公应用-新闻管理-新闻浏览", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_新闻管理_新闻浏览.png", doc,
                                  "办公应用_新闻管理_新闻浏览")
        except Exception as e:
            utils.g_logger.info("办公应用-新闻管理-新闻浏览页异常")
            sheet.append(["办公应用-新闻管理-新闻浏览", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_新闻管理_新闻浏览.png", doc,
                                  "办公应用_新闻管理_新闻浏览")

        # 新闻管理-新闻上传页
        time.sleep(5)
        self.driver.find_element(By.XPATH, "//li[@role='menuitem'][text()='新闻上传 ']").click()
        try:
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.presence_of_element_located((By.XPATH, "//div[@class='cell'][text()='标题']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-新闻管理-新闻浏览页，且指定元素存在")
            utils.g_logger.info("办公应用-新闻管理-新闻上传页显示正常")
            sheet.append(["办公应用-新闻管理-新闻上传", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_新闻管理_新闻上传.png", doc,
                                  "办公应用_新闻管理_新闻上传")
        except Exception as e:
            utils.g_logger.info("办公应用-新闻管理-新闻上传页异常")
            sheet.append(["办公应用-新闻管理-新闻上传", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_新闻管理_新闻上传.png", doc,
                                  "办公应用_新闻管理_新闻上传")"""

        # 通知公告-公告管理页
        try:
            time.sleep(5)
            self.driver.find_element(By.XPATH, "//span[text()='通知公告']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='通知公告']/parent::div/following-sibling::ul//li[text()='公告管理 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='公告名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-通知公告-公告管理页")
            utils.g_logger.info("办公应用-通知公告-公告管理页显示正常")
            sheet.append(["办公应用-通知公告-公告管理", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_通知公告_公告管理.png", doc,
                                  "办公应用_通知公告_公告管理")
        except Exception as e:
            utils.g_logger.info("办公应用-通知公告-公告管理页异常")
            sheet.append(["办公应用-通知公告-公告管理", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_通知公告_公告管理.png", doc,
                                  "办公应用_通知公告_公告管理")

        # 通知公告-公告查阅页
        time.sleep(5)
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='通知公告']/parent::div/following-sibling::ul//li[text()='公告查阅 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//div[@class='cell'][text()='公告名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开办公应用-通知公告-公告查阅页")
            utils.g_logger.info("办公应用-通知公告-公告查阅页显示正常")
            sheet.append(["办公应用-通知公告-公告查阅", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_通知公告_公告查阅.png", doc,
                                  "办公应用_通知公告_公告查阅")
        except Exception as e:
            utils.g_logger.info("办公应用-通知公告-公告查阅页异常")
            sheet.append(["办公应用-通知公告-公告查阅", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/办公应用_通知公告_公告查阅.png", doc,
                                  "办公应用_通知公告_公告查阅")

        # 关闭所有页面
        self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
        try:
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_zhibaotixi(self):
        # 植保体系-上报机构信息
        time.sleep(2)
        try:
            self.driver.find_element(By.XPATH, "//*[text()='植保体系 ']").click()
            time.sleep(8)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//label[text()='植保机构名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-植保机构信息上报页")
            utils.g_logger.info("植保体系-植保机构信息上报页显示正常")
            sheet.append(["植保体系-植保机构信息上报", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保机构信息上报.png", doc,
                                  "植保体系-植保机构信息上报")
        except Exception as e:
            utils.g_logger.info("植保体系-植保机构信息上报页异常")
            sheet.append(["植保体系-植保机构信息上报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保机构信息上报.png", doc,
                                  "植保体系-植保机构信息上报")

        # 植保体系-植保人员管理-植保人员信息上报
        try:
            self.driver.find_element(By.XPATH, "//span[text()='植保人员管理']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='植保人员管理']/parent::div/following-sibling::ul//li[text()='植保人员信息上报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='编制情况']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-植保人员管理-植保人员信息上报页")
            utils.g_logger.info("植保体系-植保人员管理-植保人员信息上报页显示正常")
            sheet.append(["植保体系-植保人员管理-植保人员信息上报", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-植保人员信息上报.png", doc,
                                  "植保体系-植保人员管理-植保人员信息上报")
        except Exception as e:
            utils.g_logger.info("植保体系-植保人员管理-植保人员信息上报页异常")
            sheet.append(["植保体系-植保人员管理-植保人员信息上报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-植保人员信息上报.png", doc,
                                  "植保体系-植保人员管理-植保人员信息上报")

        # 植保体系-植保人员管理-乡镇植保人员信息上报

        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='植保人员管理']/parent::div/following-sibling::ul//li[text()='乡镇植保人员信息上报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='所属乡镇机构名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-植保人员管理-乡镇植保人员信息上报页")
            utils.g_logger.info("植保体系-植保人员管理-乡镇植保人员信息上报页显示正常")
            sheet.append(["植保体系-植保人员管理-乡镇植保人员信息上报", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-乡镇植保人员信息上报.png", doc,
                                  "植保体系-植保人员管理-乡镇植保人员信息上报")
        except Exception as e:
            utils.g_logger.info("植保体系-植保人员管理-乡镇植保人员信息上报页异常")
            sheet.append(["植保体系-植保人员管理-乡镇植保人员信息上报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-乡镇植保人员信息上报.png", doc,
                                  "植保体系-植保人员管理-乡镇植保人员信息上报")

        # 植保体系-植保人员管理-村级植保员信息上报
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='植保人员管理']/parent::div/following-sibling::ul//li[text()='村级植保员信息上报 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='服务行政村数量（个）']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-植保人员管理-村级植保员信息上报页")
            utils.g_logger.info("植保体系-植保人员管理-村级植保员信息上报页显示正常")
            sheet.append(["植保体系-植保人员管理-村级植保员信息上报", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-村级植保员信息上报.png", doc,
                                  "植保体系-植保人员管理-村级植保员信息上报")
        except Exception as e:
            utils.g_logger.info("植保体系-植保人员管理-村级植保员信息上报页异常")
            sheet.append(["植保体系-植保人员管理-村级植保员信息上报", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-植保人员管理-村级植保员信息上报.png", doc,
                                  "植保体系-植保人员管理-村级植保员信息上报")

        # 植保体系-数据审核
        try:
            self.driver.find_element(By.XPATH, "//span[contains(@style,'margin-left')][text()='数据审核']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//div[@class='search-container']//label[text()='审核状态']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据审核页")
            utils.g_logger.info("植保体系-数据审核页显示正常")
            sheet.append(["植保体系-数据审核", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据审核.png", doc,
                                  "植保体系-数据审核")
        except Exception as e:
            utils.g_logger.info("植保体系-数据审核页异常")
            sheet.append(["植保体系-数据审核", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据审核.png", doc,
                                  "植保体系-数据审核")

        # 植保体系-数据查询-机构查询
        try:
            self.driver.find_element(By.XPATH, "//span[text()='数据查询']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='数据查询']/parent::div/following-sibling::ul//li[text()='机构查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='植保机构名称']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据查询-机构查询页")
            utils.g_logger.info("植保体系-数据查询-机构查询页显示正常")
            sheet.append(["植保体系-数据查询-机构查询", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-机构查询.png", doc,
                                  "植保体系-数据查询-机构查询")
        except Exception as e:
            utils.g_logger.info("植保体系-数据查询-机构查询页异常")
            sheet.append(["植保体系-数据查询-机构查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-机构查询.png", doc,
                                  "植保体系-数据查询-机构查询")

        # 植保体系-数据查询-人员查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='数据查询']/parent::div/following-sibling::ul//li[text()='人员查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='人员类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据查询-人员查询页")
            utils.g_logger.info("植保体系-数据查询-人员查询页显示正常")
            sheet.append(["植保体系-数据查询-人员查询", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-人员查询.png", doc,
                                  "植保体系-数据查询-人员查询")
        except Exception as e:
            utils.g_logger.info("植保体系-数据查询-人员查询页异常")
            sheet.append(["植保体系-数据查询-人员查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-人员查询.png", doc,
                                  "植保体系-数据查询-人员查询")

        # 植保体系-数据查询-通讯录查询
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='数据查询']/parent::div/following-sibling::ul//li[text()='通讯录查询 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[@class='has-gutter']//div[@class='cell'][text()='手机号码']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据查询-通讯录查询页")
            utils.g_logger.info("植保体系-数据查询-通讯录查询页显示正常")
            sheet.append(["植保体系-数据查询-通讯录查询", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-通讯录查询.png", doc,
                                  "植保体系-数据查询-通讯录查询")
        except Exception as e:
            utils.g_logger.info("植保体系-数据查询-通讯录查询页异常")
            sheet.append(["植保体系-数据查询-通讯录查询", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据查询-通讯录查询.png", doc,
                                  "植保体系-数据查询-通讯录查询")

        # 植保体系-数据统计-数据统计
        try:
            self.driver.find_element(By.XPATH, "//span[text()='数据统计']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='数据统计']/parent::div/following-sibling::ul//li[text()='数据统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[contains(@class,'has-gutter')]//div[@class='cell'][text()='机构数量（个）']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-数据统计页")
            utils.g_logger.info("植保体系-数据统计-数据统计页显示正常")
            sheet.append(["植保体系-数据统计-数据统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-数据统计.png", doc,
                                  "植保体系-数据统计-数据统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-数据统计页异常")
            sheet.append(["植保体系-数据统计-数据统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-数据统计.png", doc,
                                  "植保体系-数据统计-数据统计")

        # 植保体系-数据统计-汇总统计-机构信息统计
        try:
            self.driver.find_element(By.XPATH, "//span[text()='汇总统计']").click()
            time.sleep(2)
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='机构信息统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[contains(@class,'has-gutter')]//div[@class='cell'][text()='植保机构个数']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-机构信息统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-机构信息统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-机构信息统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-机构信息统计.png", doc,
                                  "植保体系-数据统计-汇总统计-机构信息统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-机构信息统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-机构信息统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-机构信息统计.png", doc,
                                  "植保体系-数据统计-汇总统计-机构信息统计")

        # 植保体系-数据统计-汇总统计-植保机构统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='植保机构统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//label[text()='行政区划层级']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-植保机构统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-植保机构统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-植保机构统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-植保机构统计.png", doc,
                                  "植保体系-数据统计-汇总统计-植保机构统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-植保机构统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-植保机构统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-植保机构统计.png", doc,
                                  "植保体系-数据统计-汇总统计-植保机构统计")

        # 植保体系-数据统计-汇总统计-机构类别统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='机构类别统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//label[text()='机构类型']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-机构类别统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-机构类别统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-机构类别统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-机构类别统计.png", doc,
                                  "植保体系-数据统计-汇总统计-机构类别统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-机构类别统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-机构类别统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-机构类别统计.png", doc,
                                  "植保体系-数据统计-汇总统计-机构类别统计")

        # 植保体系-数据统计-汇总统计-人员年龄统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='人员年龄统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//span[text()='人员年龄汇总（植保机构人员）：']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-人员年龄统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-人员年龄统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-人员年龄统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-人员年龄统计.png", doc,
                                  "植保体系-数据统计-汇总统计-人员年龄统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-人员年龄统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-人员年龄统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-人员年龄统计.png", doc,
                                  "植保体系-数据统计-汇总统计-人员年龄统计")

        # 植保体系-数据统计-汇总统计-人员学历统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='人员学历统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//span[text()='人员学历汇总（植保机构人员）：']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-人员学历统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-人员学历统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-人员学历统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-人员学历统计.png", doc,
                                  "植保体系-数据统计-汇总统计-人员学历统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-人员学历统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-人员学历统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-人员学历统计.png", doc,
                                  "植保体系-数据统计-汇总统计-人员学历统计")

        # 植保体系-数据统计-汇总统计-填报进度统计
        try:
            self.driver.find_element(By.XPATH,
                                     "//span[text()='汇总统计']/parent::div/following-sibling::ul//li[text()='填报进度统计 ']").click()
            time.sleep(10)
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located(
                    (By.XPATH, "//thead[contains(@class,'has-gutter')]//div[@class='cell'][text()='填报率（%）']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开植保体系-数据统计-汇总统计-填报进度统计页")
            utils.g_logger.info("植保体系-数据统计-汇总统计-填报进度统计页显示正常")
            sheet.append(["植保体系-数据统计-汇总统计-填报进度统计", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-填报进度统计.png", doc,
                                  "植保体系-数据统计-汇总统计-填报进度统计")
        except Exception as e:
            utils.g_logger.info("植保体系-数据统计-汇总统计-填报进度统计页异常")
            sheet.append(["植保体系-数据统计-汇总统计-填报进度统计", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/植保体系-数据统计-汇总统计-填报进度统计.png", doc,
                                  "植保体系-数据统计-汇总统计-填报进度统计")
        # 执行“关闭所有页面”操作
        try:
            time.sleep(2)
            self.driver.find_element(By.XPATH, "//div[@class='close-con']").click()
            time.sleep(2)
            element = WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(
                (By.XPATH, "//ul[@class='el-dropdown-menu el-popper']/li[text()='关闭所有']")))
            element.click()
            utils.g_logger.info("成功关闭所有页面")
        except Exception:
            utils.g_logger.info("关闭所有页面失败")

    @utils.retry(MAX_TRIES)
    def test_shuzituku(self):
        # 电子标本库页
        time.sleep(3)
        try:
            self.driver.find_element(By.XPATH, "//*[text()='电子标本库 ']").click()
            time.sleep(5)
            # 获取所有窗口的句柄
            windows = self.driver.window_handles
            # 切换到新窗口，通常新窗口是最后一个句柄
            self.driver.switch_to.window(windows[-1])
            element = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, "//li[text()=' 标本数据 ']")))
            unittest.TestCase.assertTrue(element is not None, "成功打开电子标本库页")
            utils.g_logger.info("电子标本库页显示正常")
            sheet.append(["电子标本库", "正常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/电子标本库.png", doc,
                                  "电子标本库")
            self.driver.close()
            self.driver.switch_to.window(windows[0])
        except Exception as e:
            utils.g_logger.info("电子标本库页异常")
            sheet.append(["电子标本库", "异常"])
            utils.page_screenshot(self.driver, "outputs/imagefiles/电子标本库.png", doc,
                                  "电子标本库")

    def export_excel(self):
        # 退出浏览器
        self.driver.quit()
        # 调整列宽
        for col in sheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 14)
            sheet.column_dimensions[column].width = adjusted_width
        wb.save('outputs/省级功能巡检.xlsx')
        # 保存截图的word文件
        doc.save('outputs/新疆截图.docx')
        utils.g_logger.info("省级系统功能巡检结束，请检查输出的巡检文件内容。")


if __name__ == '__main__':
    p = PPMSXJ()
    # 新疆省级系统页面巡检
    p.test_shouye()
    p.test_shujutianbao()
    p.test_wulianwang()
    p.test_zhiwujianyi()
    p.test_zhishiku()
    p.test_zhibaotixi()
    p.test_xitongguanli()
    p.test_shuzituku()  # 必须放在最后执行页面检查方法
    # 将所有巡检结果导出excel文件
    p.export_excel()
